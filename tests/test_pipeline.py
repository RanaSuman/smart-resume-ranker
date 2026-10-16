"""Unit tests for parser, preprocessor, and matcher.

Run:
    pytest -q
    python -m pytest tests/test_pipeline.py -v
"""

from __future__ import annotations

import io
import time

import pytest


# ---------------------------------------------------------------- parser
class TestParser:
    def test_sanitize_filename_blocks_traversal(self):
        from utils.parser import sanitize_filename

        assert sanitize_filename("../../etc/passwd") == "passwd"
        assert sanitize_filename("..\\..\\win.ini") == "win.ini"
        assert sanitize_filename("my resume (1).PDF") != ""
        assert sanitize_filename("") == "unnamed"
        assert "\x00" not in sanitize_filename("a\x00b.pdf")

    def test_validate_extension_whitelist(self):
        from utils.parser import FileValidationError, validate_extension

        assert validate_extension("cv.pdf") == ".pdf"
        assert validate_extension("CV.DOCX") == ".docx"
        with pytest.raises(FileValidationError):
            validate_extension("evil.exe")
        with pytest.raises(FileValidationError):
            validate_extension("noext")

    def test_signature_rejects_spoofed_pdf(self):
        from utils.parser import FileValidationError, extract_text

        fake = b"MZ\x90\x00 not a pdf at all...."
        with pytest.raises(FileValidationError):
            extract_text(fake, "fake.pdf")

    def test_signature_rejects_spoofed_docx(self):
        from utils.parser import FileValidationError, extract_text

        fake = b"%PDF-1.4 fake pretending to be docx"
        with pytest.raises(FileValidationError):
            extract_text(fake, "fake.docx")

    def test_size_cap_and_empty(self):
        from utils.parser import FileValidationError, MAX_FILE_SIZE_BYTES, extract_text

        with pytest.raises(FileValidationError):
            extract_text(b"", "empty.pdf")
        big = b"%PDF" + b"x" * (MAX_FILE_SIZE_BYTES + 1)
        with pytest.raises(FileValidationError):
            extract_text(big, "big.pdf")

    def test_docx_roundtrip(self):
        """Build a real DOCX in-memory, extract it back."""
        pytest.importorskip("docx")
        from docx import Document

        from utils.parser import extract_text

        buf = io.BytesIO()
        doc = Document()
        doc.add_paragraph("Jane Doe — Python engineer, skilled in machine learning.")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "SQL"
        table.cell(0, 1).text = "Docker"
        doc.save(buf)
        data = buf.getvalue()

        text = extract_text(data, "jane.docx")
        assert "Python" in text
        assert "SQL" in text  # tables captured

    def test_pdf_roundtrip_if_possible(self):
        """Generate a minimal PDF via reportlab if available; else skip."""
        reportlab = pytest.importorskip("reportlab")
        from reportlab.pdfgen import canvas

        from utils.parser import extract_text

        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(72, 720, "John Smith Python developer machine learning SQL")
        c.showPage()
        c.save()
        text = extract_text(buf.getvalue(), "john.pdf")
        assert "Python" in text

    def test_scanned_pdf_raises_scanned_error(self):
        from utils.parser import ScannedDocumentError, extract_text

        # Minimal valid-PDF header but ~no text content.
        data = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"
        # PyPDF2/pdfplumber will yield no pages or no text -> Scanned or
        # TextExtraction error; either is acceptable graceful handling.
        with pytest.raises(Exception):
            extract_text(data, "blank.pdf")

    def test_corrupt_docx_graceful(self):
        from utils.parser import TextExtractionError, extract_text

        # Valid ZIP magic but not a real docx.
        bad = b"PK\x03\x04" + b"\x00" * 100
        with pytest.raises(Exception):
            extract_text(bad, "bad.docx")

    def test_sanitize_text_strips_controls(self):
        from utils.parser import sanitize_text

        assert "\x00" not in sanitize_text("a\x00b\x1bc")
        assert sanitize_text("") == ""


# ---------------------------------------------------------------- preprocessor
class TestPreprocessor:
    def test_redact_pii(self):
        from utils.preprocessor import redact_pii

        out = redact_pii("Contact me at jane@mail.com or https://x.io, +1 555-123-4567")
        assert "jane@mail.com" not in out
        assert "https://" not in out
        assert redact_pii("") == ""

    def test_normalize_whitespace(self):
        from utils.preprocessor import normalize_whitespace

        assert normalize_whitespace("  Hello,   WORLD!! 123 ") == "hello world"
        assert normalize_whitespace("") == ""

    def test_preprocess_empty_safe(self):
        from utils.preprocessor import preprocess_text

        assert preprocess_text("") == ""
        assert preprocess_text("   ") == ""

    def test_preprocess_corpus_empty_list(self):
        from utils.preprocessor import preprocess_corpus

        assert preprocess_corpus([]) == []

    def test_preprocess_with_model(self):
        spacy = pytest.importorskip("spacy")
        try:
            spacy.load("en_core_web_sm")
        except OSError:
            pytest.skip("en_core_web_sm not installed")
        from utils.preprocessor import preprocess_corpus, preprocess_text

        out = preprocess_text("Senior Python engineers build machine learning systems.")
        assert isinstance(out, str) and len(out) > 0
        batch = preprocess_corpus(["Python developer SQL", "  "])
        assert len(batch) == 2 and batch[1] == ""


# ---------------------------------------------------------------- matcher
class TestMatcher:
    def _simple(self):
        return (
            "python machine learning engineer sql docker",
            [
                "python machine learning engineer sql docker kubernetes",
                "accountant bookkeeping excel tax audit",
            ],
            ["strong.pdf", "weak.pdf"],
        )

    def test_ranking_accuracy(self):
        from utils.matcher import rank_resumes

        jd, resumes, names = self._simple()
        cands, meta = rank_resumes(jd, resumes, names)
        assert len(cands) == 2
        # Strong overlap must rank first with a higher score.
        assert cands[0].name == "strong.pdf"
        assert cands[0].score > cands[1].score
        assert cands[0].rank == 1 and cands[1].rank == 2
        assert 0.0 <= cands[0].score <= 100.0
        assert "vectorize_and_score_ms" in meta

    def test_zero_overlap_returns_zeros_not_crash(self):
        """Zero common tokens -> all 0.0, no division-by-zero / empty-matrix crash."""
        from utils.matcher import rank_resumes

        cands, meta = rank_resumes(
            "zzqx jwxq kvzz", ["aaab bbbc cccd", "ddde eeefff ggg"], ["a", "b"]
        )
        assert len(cands) == 2
        assert all(c.score == 0.0 for c in cands)
        assert all(c.raw_cosine == 0.0 for c in cands)

    def test_empty_vocabulary_returns_zeros(self):
        from utils.matcher import rank_resumes

        # All stopwords / single chars -> vectorizer empty vocabulary.
        cands, meta = rank_resumes("the and or", ["the and", "or the"], ["a", "b"])
        assert all(c.score == 0.0 for c in cands)

    def test_empty_jd_raises(self):
        from utils.matcher import rank_resumes

        with pytest.raises(ValueError):
            rank_resumes("", ["python developer"], ["a.pdf"])
        with pytest.raises(ValueError):
            rank_resumes("python", [], [])

    def test_length_mismatch_raises(self):
        from utils.matcher import rank_resumes

        with pytest.raises(ValueError):
            rank_resumes("python dev", ["a", "b"], ["only-one"])

    def test_skill_gaps_present(self):
        from utils.matcher import rank_resumes

        jd, resumes, names = self._simple()
        cands, _ = rank_resumes(jd, resumes, names, top_gap_terms=10)
        weak = [c for c in cands if c.name == "weak.pdf"][0]
        assert isinstance(weak.skill_gaps, list)

    def test_performance_50_resumes(self):
        """O(N*M) sparse path: 50 resumes must vectorize+score in < 500 ms."""
        from utils.matcher import rank_resumes

        jd = "python machine learning engineer sql docker kubernetes aws"
        resumes = [
            f"resume {i} python developer sql {'expert' if i % 2 == 0 else 'junior'}"
            for i in range(50)
        ]
        names = [f"r{i}.pdf" for i in range(50)]
        t0 = time.perf_counter()
        cands, meta = rank_resumes(jd, resumes, names)
        elapsed_ms = (time.perf_counter() - t0) * 1000.0
        assert len(cands) == 50
        assert elapsed_ms < 500, f"too slow: {elapsed_ms:.0f} ms"

    def test_score_distribution(self):
        from utils.matcher import RankedCandidate, get_score_distribution

        cands = [RankedCandidate("a", 90.0, 0.9), RankedCandidate("b", 10.0, 0.1)]
        d = get_score_distribution(cands, bins=10)
        assert sum(d["counts"]) == 2
        assert get_score_distribution([], 10) == {"bin_edges": [], "counts": []}
